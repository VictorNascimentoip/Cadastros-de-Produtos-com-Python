from openpyxl.reader.excel import load_workbook

class ExcelManager:
    
    ARQUIVO = "Data/Produtos.xlsx"
    ABA = "Produto"
    TABELA = "Estoque"
    
    def abrir_tabela(self):
        
        wb = load_workbook(self.ARQUIVO)
        ws = wb[self.ABA]
        
        return wb, ws
    
    def atualizar_tabela(self, ws, ultima_linha):
        
        tabela = ws.tables[self.TABELA]
        tabela.ref = f"A1:D{ultima_linha}"
        
    def salvar(self, wb):
        
        wb.save(self.ARQUIVO)
    
    def proxima_linha(self, ws):
        
        linha = 2
        
        while ws.cell(row = linha, column = 1).value is not None:
            linha += 1
    
        return linha